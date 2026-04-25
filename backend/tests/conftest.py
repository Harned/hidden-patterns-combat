"""Фикстуры backend-тестов.

Изолируем SQLite и storage в tmp_path, чтобы тесты не задевали dev-базу.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from app.config import Settings, get_settings
from app.db.session import init_schema
from app.main import create_app


@pytest.fixture
def tmp_settings(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Settings:
    storage_root = tmp_path / "storage"
    db_path = tmp_path / "test.db"

    settings = Settings(
        environment="test",
        database_url=f"sqlite:///{db_path}",
        storage_root=storage_root,
        secret_key="test-secret-long-enough-for-hs256-32b",
        cookie_secure=False,
    )

    # Сбрасываем lru_cache engine-а, чтобы тест не делил движок с другим.
    from app.db import session as db_session

    db_session._build_engine.cache_clear()  # type: ignore[attr-defined]

    monkeypatch.setattr("app.config.get_settings", lambda: settings)
    monkeypatch.setattr("app.db.session.get_settings", lambda: settings)
    monkeypatch.setattr("app.main.get_settings", lambda: settings)

    init_schema(settings)
    return settings


@pytest.fixture
def client(tmp_settings: Settings) -> Iterator[TestClient]:
    app = create_app(tmp_settings)
    app.dependency_overrides[get_settings] = lambda: tmp_settings

    with TestClient(app) as c:
        yield c


@pytest.fixture
def register_verified(client: TestClient):
    """Зарегистрировать пользователя с согласиями и сразу подтвердить email.

    Удобно для тестов, где TASK_SPEC_010 не сам объект проверки. Возвращает
    тело ответа `/auth/register` (содержит `csrf_token` и id).
    """

    from datetime import UTC, datetime

    from sqlalchemy import select

    from app.db.models import User
    from app.db.session import get_sessionmaker

    def _register(email: str, password: str = "supersecret123") -> dict:
        resp = client.post(
            "/api/auth/register",
            json={
                "email": email,
                "password": password,
                "accept_terms": True,
                "accept_pdn": True,
            },
        )
        assert resp.status_code == 201, resp.text
        SessionLocal = get_sessionmaker()
        with SessionLocal() as db:
            user = db.execute(
                select(User).where(User.email == email.lower())
            ).scalar_one()
            user.email_verified_at = datetime.now(UTC)
            db.commit()
        return resp.json()

    return _register


@pytest.fixture
def sample_xlsx_bytes() -> bytes:
    """Маленький валидный Excel с ЗАП-колонкой для e2e."""

    df = pd.DataFrame(
        {
            "Спортсмен": ["Иванов", "Петров", "Иванов"],
            "Эпизод": [1, 1, 2],
            "Время, сек": [12.0, 7.5, 20.1],
            "ЗАП": ["ЗАП-Р", "ЗАП-Т", "удержание"],
        }
    )
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Общее", index=False)
    return buf.getvalue()


@pytest.fixture
def multirow_xlsx_bytes() -> bytes:
    """Excel с многострочными заголовками — минимально достаточный для
    полноценного baseline через mapping (TASK_SPEC_003)."""

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "48"
    ws.append(
        [
            "ФИО борца",
            "Технико-тактический эпизод",
            None,
            None,
            "Баллы",
            "Стойка и маневрирование самбиста (основные в эпизоде)",
            None,
            "КФВ",
            None,
            "ВУП",
        ]
    )
    ws.append(
        [
            None,
            None,
            None,
            None,
            None,
            "Правосторонняя стойка (ПС)",
            None,
            "Захваты",
            None,
            None,
        ]
    )
    ws.append(
        [
            None,
            "№ эпизода",
            "Время эпизода, с.",
            "Время паузы, с.",
            "ЗАП",
            "Вперед-влево",
            "Назад",
            "Двусторонний захват",
            "Односторонний захват",
            "Передний",
        ]
    )
    for row in (
        ["Иванов", 1, 34, 7, "ЗАП-Р", 1, 0, 1, 0, 0],
        ["Иванов", 2, 29, 8, "ЗАП-Т", 0, 1, 0, 1, 0],
        ["Петров", 1, 15, 5, "удержание", 1, 0, 1, 0, 1],
    ):
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def dense_hmm_xlsx_bytes() -> bytes:
    """Плотная фикстура для e2e-проверки HMM (TASK_SPEC_004)."""

    import random

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "48"
    ws.append(
        [
            "ФИО борца",
            "Технико-тактический эпизод",
            None,
            "Завершающие атаку приемы (n)",
            None,
            None,
            None,
        ]
    )
    ws.append([None, None, None, None, "Болевой прием", None, None])
    ws.append(
        [
            None,
            "№ эпизода",
            "Время эпизода, с.",
            "Удержание",
            "На руку",
            "На ногу",
            "ЗАП-Р",
        ]
    )

    rng = random.Random(0)
    names = ["Иванов", "Петров", "Сидоров", "Кузнецов"]
    for i in range(1, 41):
        name = rng.choice(names)
        ws.append(
            [
                name,
                i,
                rng.randint(10, 40),
                1 if rng.random() < 0.3 else 0,
                1 if rng.random() < 0.2 else 0,
                1 if rng.random() < 0.2 else 0,
                1 if rng.random() < 0.25 else 0,
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def very_dense_xlsx_bytes() -> bytes:
    """Плотная фикстура для e2e detailed-HMM (TASK_SPEC_005)."""

    import random

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "all"
    ws.append(
        [
            "ФИО борца",
            "Технико-тактический эпизод",
            None,
            "Завершающие атаку приемы (n)",
            None,
            None,
            None,
            None,
        ]
    )
    ws.append([None, None, None, None, "Болевой прием", None, None, None])
    ws.append(
        [
            None,
            "№ эпизода",
            "Время эпизода, с.",
            "Удержание",
            "На руку",
            "На ногу",
            "ЗАП-Р",
            "ЗАП-Т",
        ]
    )

    rng = random.Random(0)
    names = ["Иванов", "Петров", "Сидоров", "Кузнецов", "Смирнов", "Попов"]
    for i in range(1, 181):
        name = rng.choice(names)
        ws.append(
            [
                name,
                i,
                rng.randint(10, 40),
                1 if rng.random() < 0.35 else 0,
                1 if rng.random() < 0.25 else 0,
                1 if rng.random() < 0.2 else 0,
                1 if rng.random() < 0.25 else 0,
                1 if rng.random() < 0.2 else 0,
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def multirow_binary_xlsx_bytes() -> bytes:
    """Multi-row header с бинарной ЗАП-кодировкой для e2e-теста TASK_SPEC_003_1."""

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "48"
    ws.append(
        [
            "ФИО борца",
            "Технико-тактический эпизод",
            None,
            "Стойка и маневрирование самбиста",
            None,
            "КФВ",
            "ВУП",
            "Завершающие атаку приемы (n)",
            None,
            None,
        ]
    )
    ws.append(
        [
            None,
            None,
            None,
            "Правосторонняя стойка (ПС)",
            None,
            "Захваты",
            None,
            None,
            "Болевой прием",
            None,
        ]
    )
    ws.append(
        [
            None,
            "№ эпизода",
            "Время эпизода, с.",
            "Вперед",
            "Назад",
            "Двусторонний",
            "Передний",
            "Удержание",
            "На руку",
            "На ногу",
        ]
    )
    for row in (
        ["Иванов", 1, 34, 1, 0, 1, 0, 1, 0, 0],
        ["Иванов", 2, 29, 0, 1, 0, 1, 0, 1, 0],
        ["Петров", 1, 15, 1, 0, 1, 1, 0, 0, 1],
        ["Петров", 2, 22, 0, 0, 0, 0, 2, 0, 0],
        ["Сидоров", 1, 18, 1, 1, 0, 0, 0, 0, 0],
    ):
        ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
