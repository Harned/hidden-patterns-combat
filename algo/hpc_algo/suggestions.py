"""Предложения по предобработке данных, которые показываются пользователю
в мастере источника. Любые автоправки данных должны проходить через явное
подтверждение пользователем — модуль возвращает только список предложений
с координатами в Excel (1-based), исходное значение и источник переноса.

Поддерживаются две эвристики:

* :func:`forward_fill_athlete_suggestions` — копирование последнего непустого
  ФИО вниз в data-секции (детекции личности по эпизодам нет).
* :func:`header_merge_fill_suggestions` — материализация значений merged-ячеек
  в **зоне шапки**: «подчинённые» (slave) ячейки объединённого диапазона
  получают копию текста master, чтобы flatten-имена колонок не теряли
  верхние уровни заголовка после ручного редактирования.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from hpc_algo.mapping import read_sheet_with_header_rows


@dataclass(frozen=True)
class AthleteSuggestion:
    """Одно предложение «вписать ФИО в пустую ячейку из строки выше»."""

    row: int  # 1-based Excel row
    col: int  # 1-based Excel column
    proposed: str
    source_row: int  # 1-based Excel row, откуда взято значение
    message_ru: str


@dataclass(frozen=True)
class ForwardFillReport:
    suggestions: list[AthleteSuggestion]
    athlete_column: str | None = None
    athlete_columns_seen: list[str] = field(default_factory=list)
    warning: str | None = None


def _is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    if isinstance(value, str):
        return value.strip() == ""
    try:
        if pd.isna(value):  # numpy NaN, pd.NaT
            return True
    except (TypeError, ValueError):
        pass
    return False


def _normalize_athlete(value: object) -> str | None:
    if _is_blank(value):
        return None
    return str(value).strip() or None


def _resolve_column_index(df: pd.DataFrame, column_name: str) -> int | None:
    """Вернуть позицию колонки (0-based) с учётом возможных дубликатов имён."""

    for idx, name in enumerate(df.columns):
        if name == column_name:
            return idx
    return None


def _is_data_row_fully_empty(df: pd.DataFrame, row_index: int) -> bool:
    """Все значения в строке data-фрейма пустые (как в count_empty_rows по смыслу)."""

    if row_index < 0 or row_index >= len(df):
        return True
    row = df.iloc[row_index]
    for val in row:
        if not _is_blank(val):
            return False
    return True


def _row_has_episode_or_other_context(
    df: pd.DataFrame,
    row_index: int,
    episode_columns: list[str] | None,
) -> bool:
    """Есть ли смысловой контекст для подстановки ФИО: при известных колонках
    эпизода — достаточно непустого значения в одной из них; иначе — любой
    непустой столбец кроме athlete (уже пуст), чтобы не трогать пустую строку.
    """

    if row_index < 0 or row_index >= len(df):
        return False
    row = df.iloc[row_index]
    cols = list(df.columns)
    if episode_columns:
        for name in episode_columns:
            if name in cols and not _is_blank(row[name]):
                return True
        return False
    for c in cols:
        if not _is_blank(row[c]):
            return True
    return False


def forward_fill_athlete_suggestions(
    path: str | Path,
    sheet_name: str,
    *,
    header_rows: list[int],
    athlete_columns: list[str],
    episode_columns: list[str] | None = None,
) -> ForwardFillReport:
    """Найти ячейки ФИО, которые пусты, и предложить заполнить их значением
    последней непустой ячейки выше в той же колонке.

    Координаты ``row``/``col`` в результате — 1-based в системе Excel,
    совместимо с :mod:`hpc_algo.workbook_editor`.

    Предложения **не** создаются для полностью пустых data-строк. Если
    ``episode_columns`` задан, подстановка рассматривается только для
    строк, где в одной из этих колонок есть значение (строка «только ФИО»
    без эпизода в данных не подсвечивается).
    """

    if not athlete_columns:
        return ForwardFillReport(
            suggestions=[],
            athlete_column=None,
            athlete_columns_seen=[],
            warning=(
                "Назначьте роль «спортсмен» (athlete) хотя бы одной колонке, "
                "чтобы получить предложения по заполнению ФИО."
            ),
        )

    if not header_rows:
        header_rows = [0]
    first_data_excel_row = max(header_rows) + 2

    df = read_sheet_with_header_rows(path, sheet_name, header_rows)

    seen: list[str] = [c for c in athlete_columns if c in df.columns]
    if not seen:
        return ForwardFillReport(
            suggestions=[],
            athlete_column=None,
            athlete_columns_seen=[],
            warning=(
                "Колонки ФИО из mapping не найдены в листе при текущих "
                "header_rows. Проверьте сопоставление колонок."
            ),
        )

    target = seen[0]
    col_excel = (_resolve_column_index(df, target) or 0) + 1

    suggestions: list[AthleteSuggestion] = []
    last_value: str | None = None
    last_excel_row: int | None = None

    for offset, raw_value in enumerate(df[target].tolist()):
        excel_row = first_data_excel_row + offset
        normalized = _normalize_athlete(raw_value)
        if normalized is not None:
            last_value = normalized
            last_excel_row = excel_row
            continue
        if last_value is None or last_excel_row is None:
            continue
        if _is_data_row_fully_empty(df, offset):
            continue
        if not _row_has_episode_or_other_context(df, offset, episode_columns):
            continue
        suggestions.append(
            AthleteSuggestion(
                row=excel_row,
                col=col_excel,
                proposed=last_value,
                source_row=last_excel_row,
                message_ru=(
                    f"Перенос ФИО из строки {last_excel_row} "
                    f"(значение «{last_value}»)."
                ),
            )
        )

    warning: str | None = None
    if len(seen) > 1:
        ignored = ", ".join(f"«{c}»" for c in seen[1:])
        warning = (
            f"Используется первая колонка ФИО — «{target}». "
            f"Дополнительные колонки {ignored} пока игнорируются."
        )

    return ForwardFillReport(
        suggestions=suggestions,
        athlete_column=target,
        athlete_columns_seen=seen,
        warning=warning,
    )


@dataclass(frozen=True)
class HeaderMergeSuggestion:
    """Одно предложение «вписать значение master-ячейки в slave merged-области шапки»."""

    row: int  # 1-based Excel row (slave)
    col: int  # 1-based Excel column (slave)
    proposed: str
    source_row: int  # 1-based Excel row (master)
    source_col: int  # 1-based Excel column (master)
    message_ru: str


@dataclass(frozen=True)
class HeaderMergeReport:
    suggestions: list[HeaderMergeSuggestion]
    header_rows: list[int]
    warning: str | None = None


def _cell_text(value: object) -> str:
    if _is_blank(value):
        return ""
    return str(value).strip()


def header_merge_fill_suggestions(
    path: str | Path,
    sheet_name: str,
    *,
    header_rows: list[int],
) -> HeaderMergeReport:
    """Найти merged-диапазоны в зоне многострочной шапки и предложить заполнить
    их «подчинённые» ячейки копией значения master.

    Логика merged-only: затрагиваются **только** те ячейки, которые входят в
    объединённый диапазон, целиком расположенный в строках шапки. Остальные
    пустые ячейки шапки не трогаем (это могут быть осознанно пустые поля).

    ``header_rows`` — 0-based, как и в остальном модуле; ``row``/``col`` в
    результате — 1-based Excel-координаты.
    """

    if not header_rows:
        return HeaderMergeReport(
            suggestions=[],
            header_rows=[],
            warning="header_rows пуст — нечего заполнять.",
        )

    header_excel_rows = {r + 1 for r in header_rows}
    min_header_row = min(header_excel_rows)
    max_header_row = max(header_excel_rows)

    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            return HeaderMergeReport(
                suggestions=[],
                header_rows=sorted(header_rows),
                warning=f"Лист «{sheet_name}» не найден.",
            )
        ws = wb[sheet_name]

        suggestions: list[HeaderMergeSuggestion] = []
        for merged_range in list(ws.merged_cells.ranges):
            r1, r2 = merged_range.min_row, merged_range.max_row
            c1, c2 = merged_range.min_col, merged_range.max_col
            # Берём диапазоны, целиком лежащие в зоне шапки.
            if r1 < min_header_row or r2 > max_header_row:
                continue
            # Дополнительно: каждая строка диапазона действительно объявлена шапкой.
            if not all(r in header_excel_rows for r in range(r1, r2 + 1)):
                continue
            master_value = ws.cell(row=r1, column=c1).value
            text = _cell_text(master_value)
            if not text:
                continue
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    if r == r1 and c == c1:
                        continue
                    suggestions.append(
                        HeaderMergeSuggestion(
                            row=r,
                            col=c,
                            proposed=text,
                            source_row=r1,
                            source_col=c1,
                            message_ru=(
                                f"Перенос значения шапки из ячейки "
                                f"({r1}, {c1}) — «{text}»."
                            ),
                        )
                    )
    finally:
        wb.close()

    suggestions.sort(key=lambda s: (s.row, s.col))
    return HeaderMergeReport(
        suggestions=suggestions,
        header_rows=sorted(header_rows),
    )


__all__ = [
    "AthleteSuggestion",
    "ForwardFillReport",
    "forward_fill_athlete_suggestions",
    "HeaderMergeSuggestion",
    "HeaderMergeReport",
    "header_merge_fill_suggestions",
]
