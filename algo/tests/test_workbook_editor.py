"""Тесты редактора Excel: чтение/запись блоков и удаление пустых строк."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from hpc_algo.workbook_editor import (
    CellEdit,
    apply_cell_edits,
    count_empty_rows,
    read_grid,
    remove_empty_rows,
)


@pytest.fixture
def grid_excel(tmp_path: Path) -> Path:
    path = tmp_path / "grid.xlsx"
    df = pd.DataFrame(
        {
            "name": ["Иванов", "Петров", "Сидоров"],
            "score": [10, 20, 30],
            "note": ["a", "b", "c"],
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)
    return path


@pytest.fixture
def grid_with_empty_rows(tmp_path: Path) -> Path:
    """Лист с заголовком, тремя data-строками и двумя полностью пустыми."""

    path = tmp_path / "with_empty.xlsx"
    df = pd.DataFrame(
        [
            ["Иванов", 10, "a"],
            [None, None, None],
            ["Петров", 20, "b"],
            [None, None, None],
            ["Сидоров", 30, "c"],
        ],
        columns=["name", "score", "note"],
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)
    return path


def test_read_grid_returns_full_block(grid_excel: Path) -> None:
    fragment = read_grid(grid_excel, "Sheet1", start_row=1, start_col=1, n_rows=4, n_cols=3)

    assert fragment.total_rows == 4
    assert fragment.total_cols == 3
    assert fragment.cells[0] == ["name", "score", "note"]
    assert fragment.cells[1] == ["Иванов", 10, "a"]
    assert fragment.cells[3] == ["Сидоров", 30, "c"]


def test_read_grid_pads_beyond_sheet(grid_excel: Path) -> None:
    fragment = read_grid(grid_excel, "Sheet1", start_row=1, start_col=1, n_rows=10, n_cols=5)
    assert fragment.cells[5] == [None, None, None, None, None]
    assert fragment.cells[1] == ["Иванов", 10, "a", None, None]


def test_apply_cell_edits_writes_values(grid_excel: Path) -> None:
    n = apply_cell_edits(
        grid_excel,
        "Sheet1",
        [
            CellEdit(row=2, col=1, value="Иванова"),
            {"row": 3, "col": 2, "value": 99},
            {"row": 4, "col": 3, "value": ""},
        ],
    )
    assert n == 3
    fragment = read_grid(grid_excel, "Sheet1", start_row=1, start_col=1, n_rows=4, n_cols=3)
    assert fragment.cells[1][0] == "Иванова"
    assert fragment.cells[2][1] == 99
    assert fragment.cells[3][2] is None


def test_apply_cell_edits_merged_range_unmerges_and_writes_target_cell(
    tmp_path: Path,
) -> None:
    """MergedCell: снять merge и записать в запрошенную (row, col), не в master."""

    from openpyxl import Workbook

    path = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "H"
    ws["A2"] = "old"
    ws.merge_cells("A2:B3")
    wb.save(path)
    wb.close()
    n = apply_cell_edits(
        path, "Sheet", [CellEdit(row=3, col=2, value="new")]
    )
    assert n == 1
    wb2 = load_workbook(path, data_only=True)
    assert wb2.active["A2"].value == "old"
    assert wb2.active["B3"].value == "new"
    wb2.close()


def test_remove_empty_rows_keeps_header(grid_with_empty_rows: Path) -> None:
    deleted = remove_empty_rows(grid_with_empty_rows, "Sheet1", header_rows=[0])
    assert deleted == 2

    wb = load_workbook(grid_with_empty_rows)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    assert rows[0] == ("name", "score", "note")
    assert rows[1] == ("Иванов", 10, "a")
    assert rows[2] == ("Петров", 20, "b")
    assert rows[3] == ("Сидоров", 30, "c")
    assert len(rows) == 4


def test_remove_empty_rows_skips_when_none_empty(grid_excel: Path) -> None:
    deleted = remove_empty_rows(grid_excel, "Sheet1", header_rows=[0])
    assert deleted == 0


def test_count_empty_rows_matches_remove(grid_with_empty_rows: Path) -> None:
    """count_empty_rows должен возвращать ту же цифру, что и remove_empty_rows."""

    expected = count_empty_rows(grid_with_empty_rows, "Sheet1", header_rows=[0])
    assert expected == 2

    deleted = remove_empty_rows(grid_with_empty_rows, "Sheet1", header_rows=[0])
    assert deleted == expected

    after = count_empty_rows(grid_with_empty_rows, "Sheet1", header_rows=[0])
    assert after == 0


def test_count_empty_rows_zero_when_clean(grid_excel: Path) -> None:
    assert count_empty_rows(grid_excel, "Sheet1", header_rows=[0]) == 0


def test_count_empty_rows_unknown_sheet_raises(grid_excel: Path) -> None:
    with pytest.raises(KeyError):
        count_empty_rows(grid_excel, "MissingSheet")


def test_apply_cell_edits_unknown_sheet_raises(grid_excel: Path) -> None:
    with pytest.raises(KeyError):
        apply_cell_edits(grid_excel, "MissingSheet", [CellEdit(row=1, col=1, value="x")])


def test_read_grid_unknown_sheet_raises(grid_excel: Path) -> None:
    with pytest.raises(KeyError):
        read_grid(grid_excel, "MissingSheet")


def test_remove_empty_rows_does_not_touch_multirow_header(tmp_path: Path) -> None:
    """Если заголовок занимает несколько строк, удаление data не сдвигает шапку."""

    from openpyxl import Workbook

    path = tmp_path / "multirow.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["group", "group"])
    ws.append(["a", "b"])
    ws.append(["x", 1])
    ws.append([None, None])
    ws.append(["y", 2])
    wb.save(path)
    wb.close()

    deleted = remove_empty_rows(path, "Sheet1", header_rows=[0, 1])
    assert deleted == 1

    wb = load_workbook(path)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    assert rows[0] == ("group", "group")
    assert rows[1] == ("a", "b")
    assert rows[2] == ("x", 1)
    assert rows[3] == ("y", 2)
    assert len(rows) == 4
